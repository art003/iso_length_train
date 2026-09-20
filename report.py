from __future__ import annotations

import csv
import html
import json
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


@dataclass
class SheetResult:
    sheet_no: int
    line_id: str
    length_mm: int
    length_m: float
    formula: str
    status: str
    remarks: str
    included: list[dict]
    excluded: list[dict]
    ambiguous: list[dict]
    llm_raw: dict
    ai_calls: int
    markup_name: str = ""


def decide_status(ambiguous: list, notes: list[str], included: list) -> str:
    if not included:
        return "error"
    if ambiguous:
        return "review"
    return "ok"


def build_formula(included: list[dict], length_mm: int) -> str:
    if not included:
        return "нет включённых размеров"
    parts = [str(x["value_mm"]) for x in included]
    return " + ".join(parts) + f" = {length_mm} мм"


def write_csv(path: Path, rows: list[SheetResult]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(
            [
                "номер_листа",
                "обозначение_линии",
                "длина_мм",
                "длина_м",
                "формула",
                "статус",
                "замечания",
            ]
        )
        for r in rows:
            w.writerow(
                [
                    r.sheet_no,
                    r.line_id,
                    r.length_mm,
                    f"{r.length_m:.3f}".replace(".", ","),
                    r.formula,
                    r.status,
                    r.remarks,
                ]
            )
        total_mm = sum(r.length_mm for r in rows)
        w.writerow(["", "итого", total_mm, f"{total_mm/1000:.3f}".replace(".", ","), "", "", ""])


def write_xlsx(path: Path, rows: list[SheetResult], meta: dict | None = None) -> None:
    meta = meta or {}
    wb = Workbook()
    thin = Border(
        left=Side(style="thin", color="E4E8EE"),
        right=Side(style="thin", color="E4E8EE"),
        top=Side(style="thin", color="E4E8EE"),
        bottom=Side(style="thin", color="E4E8EE"),
    )
    head_fill = PatternFill("solid", fgColor="004C97")
    head_font = Font(color="FFFFFF", bold=True)
    total_mm = sum(r.length_mm for r in rows)

    ws0 = wb.active
    ws0.title = "сводка"
    info = [
        ("PDF", meta.get("pdf") or ""),
        ("Папка", meta.get("folder") or path.parent.name),
        ("Когда", meta.get("updated") or ""),
        ("Модель", f"{meta.get('provider') or ''} / {meta.get('model') or ''}".strip(" /")),
        ("Листов", meta.get("pages") or len(rows)),
        ("Сумма, мм", total_mm),
        ("Сумма, м", round(total_mm / 1000.0, 3)),
        ("ok", meta.get("ok") if "ok" in meta else sum(1 for r in rows if r.status == "ok")),
        ("review", meta.get("review") if "review" in meta else sum(1 for r in rows if r.status == "review")),
    ]
    ws0["A1"] = "Ведомость длин по изометриям"
    ws0["A1"].font = Font(bold=True, size=14, color="004C97")
    ws0.merge_cells("A1:B1")
    for i, (k, v) in enumerate(info, start=3):
        ws0.cell(i, 1, k).font = Font(color="6B7785")
        ws0.cell(i, 2, v)
    ws0.column_dimensions["A"].width = 16
    ws0.column_dimensions["B"].width = 52
    ws0["A13"] = "Картинки листов — в этой же папке, markup. Метки D1, D2 — overlays."
    ws0["A13"].font = Font(color="6B7785", italic=True)

    ws = wb.create_sheet("длины")
    headers = [
        "номер_листа",
        "обозначение_линии",
        "длина_мм",
        "длина_м",
        "формула",
        "статус",
        "замечания",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(wrap_text=True)
    for r in rows:
        ws.append([r.sheet_no, r.line_id, r.length_mm, r.length_m, r.formula, r.status, r.remarks])
    last = len(rows) + 2
    ws.cell(last, 2, "итого").font = Font(bold=True)
    ws.cell(last, 3, total_mm).font = Font(bold=True)
    ws.cell(last, 4, round(total_mm / 1000.0, 3)).font = Font(bold=True)
    for row in ws.iter_rows(min_row=1, max_row=last, max_col=7):
        for cell in row:
            cell.border = thin
    widths = [14, 20, 14, 12, 72, 12, 56]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:G{len(rows)+1}"
    ws.freeze_panes = "A2"
    wb.save(path)


def write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def write_readme(path: Path, meta: dict, rows: list[SheetResult]) -> None:
    total_mm = sum(r.length_mm for r in rows)
    lines = [
        f"PDF: {meta.get('pdf') or ''}",
        f"Посчитано: {meta.get('updated') or ''}",
        f"Листов: {len(rows)}",
        f"Сумма: {total_mm} мм = {total_mm/1000:.3f} м",
        f"ok: {sum(1 for r in rows if r.status == 'ok')}, review: {sum(1 for r in rows if r.status == 'review')}",
        "",
        "В этой папке:",
        "lengths.xlsx — таблица, открой в Excel (листы «сводка» и «длины»)",
        "report.html — то же в браузере, с картинками листов",
        "markup — чертёж, зелёным что вошло в сумму",
        "overlays — метки D1, D2, по ним проверять",
        "llm_raw — что решила модель, можно править руками",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_html_report(path: Path, meta: dict, rows: list[SheetResult]) -> None:
    total_mm = sum(r.length_mm for r in rows)
    trs = []
    for r in rows:
        img = html.escape(r.markup_name) if r.markup_name else ""
        img_cell = f'<a href="markup/{img}"><img src="markup/{img}" alt=""></a>' if img else ""
        trs.append(
            "<tr>"
            f"<td>{r.sheet_no}</td>"
            f"<td>{html.escape(r.line_id)}</td>"
            f"<td>{r.length_mm}</td>"
            f"<td>{r.length_m:.3f}</td>"
            f"<td class='{html.escape(r.status)}'>{html.escape(r.status)}</td>"
            f"<td>{html.escape(r.formula)}</td>"
            f"<td>{img_cell}</td>"
            "</tr>"
        )
    body = f"""<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="utf-8">
<title>{html.escape(str(meta.get('pdf') or 'длины'))}</title>
<style>
body {{ font-family: Arial, sans-serif; color: #2c3844; margin: 24px; background: #f3f3f3; }}
.card {{ background: #fff; padding: 20px 24px; max-width: 1100px; }}
h1 {{ color: #004C97; font-size: 22px; margin: 0 0 8px; }}
.meta {{ color: #6b7785; font-size: 14px; margin-bottom: 16px; }}
table {{ border-collapse: collapse; width: 100%; font-size: 14px; }}
th, td {{ border-bottom: 1px solid #e4e8ee; padding: 8px; text-align: left; vertical-align: middle; }}
th {{ color: #6b7785; font-size: 12px; }}
.ok {{ color: #1f7a3a; font-weight: 700; }}
.review {{ color: #b36b00; font-weight: 700; }}
img {{ height: 64px; border: 1px solid #e4e8ee; }}
.total {{ font-weight: 700; }}
</style>
</head>
<body>
<div class="card">
<h1>Длины по изометриям</h1>
<p class="meta">
PDF: {html.escape(str(meta.get('pdf') or ''))}<br>
{html.escape(str(meta.get('updated') or ''))}
 · листов {len(rows)}
 · сумма {total_mm} мм ({total_mm/1000:.3f} м)
 · ok {sum(1 for r in rows if r.status == 'ok')}
 · review {sum(1 for r in rows if r.status == 'review')}
</p>
<table>
<thead><tr><th>лист</th><th>линия</th><th>мм</th><th>м</th><th>статус</th><th>формула</th><th>лист</th></tr></thead>
<tbody>
{''.join(trs)}
<tr class="total"><td></td><td>итого</td><td>{total_mm}</td><td>{total_mm/1000:.3f}</td><td></td><td></td><td></td></tr>
</tbody>
</table>
<p class="meta">Картинки лежат рядом: markup — что вошло в сумму, overlays — метки D1, D2.</p>
</div>
</body>
</html>
"""
    path.write_text(body, encoding="utf-8")


def metrics_text(m: dict) -> str:
    inn = (m.get("tariff_usd_per_1m_tokens") or {}).get("input", 0)
    out = (m.get("tariff_usd_per_1m_tokens") or {}).get("output", 0)
    lines = [
        f"PDF: {m.get('pdf') or ''}",
        f"Папка: {m.get('job') or ''}",
        f"Сумма: {m.get('total_mm')} мм ({m.get('total_m')} м)",
        f"ok / review: {m.get('ok')} / {m.get('review')}",
        f"Модель: {m.get('provider')} / {m.get('model')}",
        f"OCR: {m.get('ocr')}",
        f"Листов: {m.get('pages')}",
        f"Время, с: {m.get('total_time_s')}",
        f"Обращений к API: {m.get('ai_calls_total')} (повторы: {m.get('ai_retries')})",
        f"Токены: {m.get('input_tokens')} / {m.get('output_tokens')}",
        f"Тариф {m.get('tariff_date')}: ${inn} in / ${out} out за 1M",
        f"Стоимость: ${m.get('cost_total_usd')}",
        f"На лист: ${m.get('cost_per_page_usd')}",
        f"Формула: {m.get('cost_formula')}",
        m.get("note") or "",
    ]
    live = m.get("live_api") or {}
    if live and m.get("provider") == "без API":
        lines.append(
            "Прошлый живой прогон: "
            f"{live.get('provider')} / {live.get('model')}, "
            f"вызовов {live.get('ai_calls_total')}, "
            f"{live.get('total_time_s')} с, "
            f"токены {live.get('input_tokens')} / {live.get('output_tokens')}"
        )
    return "\n".join(lines)
